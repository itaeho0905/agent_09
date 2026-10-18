"""JSON 사양(spec)으로 .pptx 발표자료를 만든다.

사용법:
    python build_pptx.py <spec.json> <output.pptx>   발표자료 생성
    python build_pptx.py --check <file.pptx>         생성된 파일 검증

출력 파일이 이미 있으면 덮어쓰지 않고 오류로 끝낸다(버전 번호를 올려 다시 실행).
spec 형식은 SKILL.md의 "spec 형식" 절을 본다.
"""

import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

DEFAULT_FONT = "Noto Sans KR"
THEMES = {
    "navy": {
        "primary": "1F3864",   # 제목 띠, 표 머리글
        "accent": "C55A11",    # 강조선, 글머리표
        "text": "222222",
        "muted": "6B7280",
        "light": "EEF2F7",     # 표 줄무늬, 2단 상자 배경
        "subtitle": "DDE3EC",  # 표지 부제(primary 배경 위)
    },
    "pink": {
        "primary": "AD1457",
        "accent": "F06292",
        "text": "3A2A30",
        "muted": "8A6B78",
        "light": "FCE4EC",
        "subtitle": "F8D7E3",
    },
}
DEFAULT_THEME = THEMES["navy"]
CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "column_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "area": XL_CHART_TYPE.AREA,
}
MAX_BULLETS = 7
MAX_BULLET_CHARS = 60

warnings = []


# ---------- 공통 도우미 ----------

def rgb(hex_str):
    return RGBColor.from_string(hex_str.lstrip("#"))


def resolve_theme(value):
    """spec의 theme 값을 색 사전으로 바꾼다.
    "pink"처럼 프리셋 이름, {"preset": "pink", "accent": ...}처럼 프리셋+덮어쓰기, 색 사전 모두 받는다."""
    if value is None:
        value = {}
    if isinstance(value, str):
        value = {"preset": value}
    value = dict(value)
    preset = value.pop("preset", "navy")
    if preset not in THEMES:
        sys.exit(f"오류: theme '{preset}'는 없음. 가능한 값: {', '.join(THEMES)}")
    return {**THEMES[preset], **value}


def apply_font(font, name, size=None, bold=None, color=None):
    """글꼴을 지정한다. 한글이 지정한 글꼴로 나오려면 latin과 ea(동아시아) 둘 다 지정해야 한다."""
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = rgb(color)
    font.name = name  # a:latin
    rPr = font._rPr
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
        rPr.find(qn("a:latin")).addnext(ea)
    ea.set("typeface", name)


def add_text(slide, left, top, width, height, text, ctx, size=18, bold=False,
             color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        apply_font(run.font, ctx["font"], size, bold, color or ctx["theme"]["text"])
    return box


def set_bullet(paragraph, char, color, level):
    """문단에 글머리 기호를 단다(python-pptx 공개 API에 없어 XML로 처리)."""
    pPr = paragraph._p.get_or_add_pPr()
    indent = Emu(Inches(0.3))
    pPr.set("marL", str(int(indent * (level + 1))))
    pPr.set("indent", str(-int(indent)))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color)
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def fill_bullets(tf, items, ctx, size=20):
    """items: 문자열 또는 {"text": ..., "level": 0~2} 목록."""
    tf.word_wrap = True
    if len(items) > MAX_BULLETS:
        warnings.append(f"{ctx['where']}: 글머리표 {len(items)}개 — {MAX_BULLETS}개 이하로 나누는 것이 좋음")
    for i, item in enumerate(items):
        if isinstance(item, str):
            item = {"text": item, "level": 0}
        level = int(item.get("level", 0))
        text = item["text"]
        if len(text) > MAX_BULLET_CHARS:
            warnings.append(f"{ctx['where']}: 글머리표가 {len(text)}자 — 넘칠 수 있음: {text[:20]}…")
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10 if level == 0 else 4)
        set_bullet(p, "■" if level == 0 else "–", ctx["theme"]["accent"], level)
        run = p.add_run()
        run.text = text
        apply_font(run.font, ctx["font"], size - 3 * level, False, ctx["theme"]["text"])


def add_notes(slide, spec, ctx):
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]


def content_frame(prs, spec, ctx):
    """제목 띠 + 쪽 번호가 있는 본문 슬라이드를 만들고, 본문 영역(left, top, width, height)을 돌려준다."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃
    W, H = prs.slide_width, prs.slide_height
    theme = ctx["theme"]
    add_text(slide, Inches(0.6), Inches(0.35), W - Inches(1.2), Inches(0.9),
             spec.get("title", ""), ctx, size=28, bold=True, color=theme["primary"],
             anchor=MSO_ANCHOR.MIDDLE)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.25), Inches(1.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(theme["accent"])
    bar.line.fill.background()
    add_text(slide, W - Inches(1.6), H - Inches(0.55), Inches(1.0), Inches(0.35),
             str(len(prs.slides)), ctx, size=11, color=theme["muted"], align=PP_ALIGN.RIGHT)
    if ctx.get("footer"):
        add_text(slide, Inches(0.6), H - Inches(0.55), Inches(6), Inches(0.35),
                 ctx["footer"], ctx, size=11, color=theme["muted"])
    add_notes(slide, spec, ctx)
    return slide, (Inches(0.6), Inches(1.6), W - Inches(1.2), H - Inches(2.4))


# ---------- 슬라이드 종류 ----------

def slide_title(prs, spec, ctx, closing=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W, H = prs.slide_width, prs.slide_height
    theme = ctx["theme"]
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(theme["primary"])
    add_text(slide, Inches(0.9), Inches(2.3), W - Inches(1.8), Inches(1.5),
             spec.get("title", ""), ctx, size=40 if not closing else 36, bold=True,
             color="FFFFFF", anchor=MSO_ANCHOR.BOTTOM)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.95), Inches(1.6), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(theme["accent"])
    line.line.fill.background()
    if spec.get("subtitle"):
        add_text(slide, Inches(0.9), Inches(4.2), W - Inches(1.8), Inches(1.4),
                 spec["subtitle"], ctx, size=20, color=theme["subtitle"])
    add_notes(slide, spec, ctx)


def slide_section(prs, spec, ctx):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    W = prs.slide_width
    theme = ctx["theme"]
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(theme["light"])
    if spec.get("label"):
        add_text(slide, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(0.6),
                 spec["label"], ctx, size=18, bold=True, color=theme["accent"])
    add_text(slide, Inches(0.9), Inches(3.0), W - Inches(1.8), Inches(1.2),
             spec.get("title", ""), ctx, size=36, bold=True, color=theme["primary"])
    add_notes(slide, spec, ctx)


def slide_bullets(prs, spec, ctx):
    slide, (l, t, w, h) = content_frame(prs, spec, ctx)
    box = slide.shapes.add_textbox(l, t, w, h)
    fill_bullets(box.text_frame, spec.get("bullets", []), ctx, size=spec.get("font_size", 20))


def slide_two_column(prs, spec, ctx):
    slide, (l, t, w, h) = content_frame(prs, spec, ctx)
    gap = Inches(0.4)
    col_w = int((w - gap) / 2)
    for i, key in enumerate(("left", "right")):
        col = spec.get(key, {})
        x = l + i * (col_w + gap)
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, t, col_w, h)
        panel.fill.solid()
        panel.fill.fore_color.rgb = rgb(ctx["theme"]["light"])
        panel.line.fill.background()
        add_text(slide, x + Inches(0.3), t + Inches(0.2), col_w - Inches(0.6), Inches(0.6),
                 col.get("heading", ""), ctx, size=22, bold=True, color=ctx["theme"]["primary"])
        box = slide.shapes.add_textbox(x + Inches(0.3), t + Inches(0.95), col_w - Inches(0.6), h - Inches(1.1))
        fill_bullets(box.text_frame, col.get("bullets", []), ctx, size=spec.get("font_size", 18))


def slide_table(prs, spec, ctx):
    slide, (l, t, w, h) = content_frame(prs, spec, ctx)
    columns, rows = spec["columns"], spec["rows"]
    size = spec.get("font_size", 16 if len(rows) <= 6 else 13)
    row_h = Inches(0.5 if size >= 16 else 0.42)
    if (len(rows) + 1) * row_h > h:
        warnings.append(f"{ctx['where']}: 표 행이 {len(rows)}개 — 슬라이드를 넘칠 수 있음")
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), l, t, w, row_h * (len(rows) + 1))
    table = shape.table
    widths = spec.get("col_widths")  # 비율 목록, 예: [1, 3]
    if widths:
        total = sum(widths)
        for i, ratio in enumerate(widths):
            table.columns[i].width = int(w * ratio / total)
    theme = ctx["theme"]
    for r, values in enumerate([columns] + rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.text = ""
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(theme["primary"] if r == 0 else ("FFFFFF" if r % 2 else theme["light"]))
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(value)
            apply_font(run.font, ctx["font"], size, r == 0, "FFFFFF" if r == 0 else theme["text"])


def slide_chart(prs, spec, ctx):
    slide, (l, t, w, h) = content_frame(prs, spec, ctx)
    kind = spec.get("chart_type", "column")
    if kind not in CHART_TYPES:
        raise ValueError(f"chart_type '{kind}'는 지원하지 않음. 가능한 값: {', '.join(CHART_TYPES)}")
    data = CategoryChartData()
    data.categories = spec["categories"]
    for s in spec["series"]:
        data.add_series(s["name"], s["values"])
    chart = slide.shapes.add_chart(CHART_TYPES[kind], l, t, w, h, data).chart
    apply_font(chart.font, ctx["font"], 14, None, ctx["theme"]["text"])
    multi = len(spec["series"]) > 1
    chart.has_legend = multi or kind in ("pie", "doughnut")
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    plot = chart.plots[0]
    if spec.get("data_labels", True):
        plot.has_data_labels = True
        labels = plot.data_labels
        apply_font(labels.font, ctx["font"], 12)
        if kind in ("pie", "doughnut"):
            labels.show_percentage = True
            labels.show_value = False
            labels.number_format = "0%"
            labels.number_format_is_linked = False
    if kind in ("pie", "doughnut"):
        plot.vary_by_categories = True
    elif not multi:
        fmt = plot.series[0].format
        if kind == "line":
            fmt.line.color.rgb = rgb(ctx["theme"]["primary"])
        else:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = rgb(ctx["theme"]["primary"])


def slide_image(prs, spec, ctx):
    slide, (l, t, w, h) = content_frame(prs, spec, ctx)
    from PIL import Image
    path = spec["path"]
    if not Path(path).is_file():
        raise FileNotFoundError(f"그림 파일이 없음: {path}")
    caption_h = Inches(0.5) if spec.get("caption") else 0
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, (h - caption_h) / ih)
    pw, ph = int(iw * scale), int(ih * scale)
    slide.shapes.add_picture(path, l + int((w - pw) / 2), t, pw, ph)
    if spec.get("caption"):
        add_text(slide, l, t + ph + Inches(0.1), w, Inches(0.4), spec["caption"], ctx,
                 size=13, color=ctx["theme"]["muted"], align=PP_ALIGN.CENTER)


BUILDERS = {
    "title": slide_title,
    "section": slide_section,
    "bullets": slide_bullets,
    "two_column": slide_two_column,
    "table": slide_table,
    "chart": slide_chart,
    "image": slide_image,
    "closing": lambda prs, spec, ctx: slide_title(prs, spec, ctx, closing=True),
}


# ---------- 생성과 검증 ----------

def build(spec_path, out_path):
    out = Path(out_path)
    if out.exists():
        sys.exit(f"오류: {out}이(가) 이미 있음. 덮어쓰지 않으므로 버전 번호를 올려 다시 실행할 것.")
    spec = json.loads(Path(spec_path).read_text(encoding="utf-8-sig"))  # PowerShell이 붙이는 BOM 허용

    prs = Presentation()
    if spec.get("size", "16:9") == "16:9":
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    ctx = {
        "font": spec.get("font", DEFAULT_FONT),
        "theme": resolve_theme(spec.get("theme")),
        "footer": spec.get("footer", ""),
    }
    for i, s in enumerate(spec["slides"], start=1):
        kind = s.get("type", "bullets")
        if kind not in BUILDERS:
            sys.exit(f"오류: {i}번 슬라이드의 type '{kind}'는 지원하지 않음. 가능한 값: {', '.join(BUILDERS)}")
        ctx["where"] = f"{i}번 슬라이드"
        BUILDERS[kind](prs, s, ctx)

    props = prs.core_properties
    props.title = spec.get("title", "")
    props.author = spec.get("author", "")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"저장: {out} (슬라이드 {len(prs.slides)}장)")
    for w in warnings:
        print(f"경고: {w}")


def check(path):
    prs = Presentation(path)
    print(f"파일: {path}")
    print(f"크기: {prs.slide_width.inches:.2f} x {prs.slide_height.inches:.2f} in, 슬라이드 {len(prs.slides)}장")
    fonts = set()
    W, H = prs.slide_width, prs.slide_height
    for i, slide in enumerate(prs.slides, start=1):
        first, first_size = "", 0
        kinds = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                # 글자가 가장 큰 텍스트를 슬라이드 제목으로 본다
                sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
                if sizes and max(sizes) > first_size:
                    first, first_size = shape.text_frame.text.strip().split("\n")[0], max(sizes)
            kinds.append("표" if shape.has_table else "차트" if shape.has_chart else
                         "그림" if shape.shape_type == 13 else None)
            if shape.left + shape.width > W + 1 or shape.top + shape.height > H + 1:
                print(f"  경고: {i}번 슬라이드 '{shape.name}'이(가) 슬라이드 밖으로 나감")
        for el in slide.shapes._spTree.iter(qn("a:latin"), qn("a:ea")):
            fonts.add(el.get("typeface"))
        extra = ", ".join(k for k in kinds if k)
        notes = " [노트]" if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip() else ""
        print(f"  {i:>2}. {first[:40]}{' (' + extra + ')' if extra else ''}{notes}")
    print(f"사용 글꼴: {', '.join(sorted(f for f in fonts if f)) or '(지정 없음)'}")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--check":
        check(sys.argv[2])
    elif len(sys.argv) == 3:
        build(sys.argv[1], sys.argv[2])
    else:
        sys.exit(__doc__)
